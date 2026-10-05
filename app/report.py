# app/report.py
import os
from copy import deepcopy
from datetime import datetime
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

from app.config import REPORTS_DIR

TEMPLATE_PATH = "app/templates/report_template_v3.docx"
OUTPUT_DIR = REPORTS_DIR
BLOCKS_PER_PAGE = 2

TITLE_TO_PREFIX = {
    "Фотоотчет по общестроительным работам": "Фотоотчет по ОСР",
    "Фотоотчет по механическим системам": "Фотоотчет по МС",
    "Фотоотчет по электрическим системам": "Фотоотчет по ЭОМ",
    "Фотоотчет по слаботочным системам": "Фотоотчет по СС",
}

TITLE_TO_PHOTO_WIDTH = {
    "Фотоотчет по общестроительным работам": Mm(70),
    "Фотоотчет по механическим системам": Mm(115),
    "Фотоотчет по электрическим системам": Mm(90),
    "Фотоотчет по слаботочным системам": Mm(70),
}


def _resolve_prefix(title: str) -> str:
    t = (title or "").strip()
    for key, prefix in TITLE_TO_PREFIX.items():
        if t.startswith(key):
            return prefix
    return "Фотоотчет"


def _resolve_photo_width(title: str) -> Mm:
    t = (title or "").strip()
    for key, width in TITLE_TO_PHOTO_WIDTH.items():
        if t.startswith(key):
            return width
    return Mm(70)


def generate_report(report_data: dict) -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    doc = Document(TEMPLATE_PATH)

    for p in doc.paragraphs:
        for run in p.runs:
            if "{{ title }}" in run.text:
                run.text = run.text.replace("{{ title }}", report_data["title"])
            if "{{ report_date }}" in run.text:
                run.text = run.text.replace("{{ report_date }}", report_data["report_date"])

    table = doc.tables[0]
    items = report_data["items"]
    photo_width = _resolve_photo_width(report_data["title"])

    if not items:
        table._tbl.remove(table.rows[0]._tr)
    else:
        _fill_block_row(table.rows[0], items[0], photo_width)
        for idx, item in enumerate(items[1:], start=1):
            new_row = _copy_row(table, table.rows[0])
            _fill_block_row(new_row, item, photo_width)
            if idx % BLOCKS_PER_PAGE == 0 and idx != len(items) - 1:
                _add_page_break_in_row(new_row)

    prefix = _resolve_prefix(report_data["title"])
    filename = f"{prefix} от {report_data['report_date']}.docx"
    output_path = os.path.join(OUTPUT_DIR, filename)
    doc.save(output_path)
    return output_path


def convert_to_pdf(docx_path: str) -> str:
    """Конвертирует DOCX в PDF. Windows — через Word, Linux — через LibreOffice."""
    import platform
    import subprocess

    pdf_path = docx_path.replace(".docx", ".pdf")

    if platform.system() == "Windows":
        from docx2pdf import convert
        convert(docx_path, pdf_path)
    else:
        subprocess.run(
            [
                "libreoffice", "--headless",
                "--convert-to", "pdf",
                "--outdir", os.path.dirname(docx_path),
                docx_path,
            ],
            check=True,
        )

    return pdf_path


def _copy_row(table, src_row):
    new_tr = deepcopy(src_row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def _add_page_break_in_row(row):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    p = row.cells[0].add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def _set_paragraph(paragraph, text, bold=False, space_after_pt=0):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(paragraph.runs):
        r.text = ""
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.bold = bold
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after_pt)


def _fill_block_row(row, item, photo_width):
    left_cell = row.cells[0]
    right_cell = row.cells[1]

    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    while len(left_cell.paragraphs) < 3:
        left_cell.add_paragraph()

    # 1-й — объект (жирный), отступ 6 пт снизу
    _set_paragraph(left_cell.paragraphs[0],
                   item.get("object_name", ""),
                   bold=True, space_after_pt=6)
    # 2-й — исполнитель (жирный), отступ 6 пт снизу
    _set_paragraph(left_cell.paragraphs[1],
                   item.get("executor", ""),
                   bold=True, space_after_pt=6)
    # 3-й — описание (обычный), без отступа
    _set_paragraph(left_cell.paragraphs[2],
                   item.get("description", ""),
                   bold=False, space_after_pt=0)

    # Правая ячейка — фото (без сжатия, оригинал)
    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(rp.runs):
        r.text = ""

    photo_path = item.get("photo_path")
    if photo_path and os.path.exists(photo_path):
        try:
            run = rp.add_run()
            run.add_picture(photo_path, width=photo_width)
        except Exception as e:
            print(f"Не удалось вставить фото {photo_path}: {e}")