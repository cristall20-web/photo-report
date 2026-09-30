# make_template_v2.py
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_style(doc, font_name="Arial", size=11):
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:cs"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)


doc = Document()
set_style(doc, "Arial", 11)

# Заголовок
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("{{ title }} от {{ report_date }}")
run.font.name = "Arial"
run.font.size = Pt(11)
run.bold = True

doc.add_paragraph()

# Таблица 2x2
table = doc.add_table(rows=2, cols=2)
table.style = "Table Grid"

# Строка 1, левая ячейка — всё в одном run, через \n
left1 = table.cell(0, 0)
p1 = left1.paragraphs[0]
p1.text = ""
run1 = p1.add_run(
    "{%tr for item in items %}\n"
    "{{ item.object_name }}\n"
    "{{ item.executor }}\n"
    "{{ item.description }}"
)
run1.font.name = "Arial"
run1.font.size = Pt(11)

# Строка 1, правая ячейка — фото
right1 = table.cell(0, 1)
r1 = right1.paragraphs[0].add_run("{{ item.photo }}")
r1.font.name = "Arial"
r1.font.size = Pt(11)

# Строка 2 — закрытие цикла
r2 = table.cell(1, 0).paragraphs[0].add_run("{%tr endfor %}")
r2.font.name = "Arial"
r2.font.size = Pt(11)

# Применяем Arial 11 ко всем ячейкам
for row in table.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for r in paragraph.runs:
                r.font.name = "Arial"
                r.font.size = Pt(11)

doc.save("app/templates/report_template.docx")
print("OK: шаблон v2 создан")