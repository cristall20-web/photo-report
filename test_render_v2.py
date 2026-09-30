# test_render_v2.py
from docxtpl import DocxTemplate
from docx import Document

t = DocxTemplate("app/templates/report_template.docx")
t.render({
    "title": "Фотоотчет по общестроительным работам",
    "report_date": "29.09.2026",
    "items": [
        {"object_name": "М7", "executor": "МК Рус", "description": "Монтаж стеновых сэндвич-панелей", "photo": None},
        {"object_name": "АБК", "executor": "КМТ", "description": "Устройство перегородок из ГКЛ", "photo": None},
    ],
})
t.save("reports/TEST_TABLE.docx")

d = Document("reports/TEST_TABLE.docx")
print("Таблиц:", len(d.tables))
if d.tables:
    for i, row in enumerate(d.tables[0].rows):
        for j, cell in enumerate(row.cells):
            print(f"[{i},{j}] = {repr(cell.text)}")